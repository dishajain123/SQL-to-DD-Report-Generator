"""Regression tests for platform-entity rewriting and T-SQL CATCH handling."""
from pathlib import Path

from app.derivation.dd_generation_engine import (
    _assignment_sites,
    _compose_simple_assignment_expression,
    _finalize_platform_expression,
    _infer_assignment_role,
    _interpret_llm_output,
    undeterminable_exception_sites,
)
from app.grammar.validator import validate_expression
from app.models.core import DerivationOption, Dialect
from app.parsing.dialect import detect_dialect
from app.parsing.object_splitter import split_objects
from app.parsing.structural_analysis import analyze_object
from app.report.dd_export import COLUMNS, export_dd_rows_excel, read_existing_dd_excel
from app.utils.sql_aliases import (
    collect_table_aliases,
    resolve_aliases_in_expression,
    rewrite_expression_to_platform_entities,
)
from datetime import date
from app.models.core import ColumnType, DDRow, DDStatus


def test_rewrite_schema_table_to_platform_entity():
    expr = 'IF("PRO"."AccountCal"."DaysPastDue" <= 90)THEN("STANDARD")ELSE("LOSS")'
    rewritten = rewrite_expression_to_platform_entities(
        expr,
        entity_name="FCT_NPA_PRODUCT",
        entity_name_map={"AccountCal": "FCT_NPA_PRODUCT"},
        alias_to_parts={"A": ("PRO", "AccountCal")},
    )
    assert '"PRO"."AccountCal"' not in rewritten
    assert '"FCT_NPA_PRODUCT"."DaysPastDue"' in rewritten
    assert validate_expression(rewritten).valid


def test_sample01_assetclass_uses_entity_not_alias_or_schema():
    sql = Path("samples/sql/01_NPA_Classification_Simple.sql").read_text()
    obj = split_objects(sql, "01.sql", detect_dialect(sql))[0]
    info = analyze_object(obj)
    sites = _assignment_sites(info, "AssetClass")
    assert len(sites) == 1
    assert sites[0].kind == "UPDATE"
    assert "BEGIN" not in sites[0].raw_sql.upper().split("UPDATE", 1)[0]

    composed = _compose_simple_assignment_expression(sites, "FCT_NPA_PRODUCT", "AssetClass")
    assert composed
    alias_map = collect_table_aliases(obj.raw_sql, obj.dialect)
    finalized = _finalize_platform_expression(
        composed,
        entity_name="FCT_NPA_PRODUCT",
        entity_name_map={"AccountCal": "FCT_NPA_PRODUCT"},
        alias_resolution_inventory=alias_map,
        source_sql=obj.raw_sql,
    )
    assert '"A"' not in finalized
    assert '"PRO"' not in finalized
    assert '"FCT_NPA_PRODUCT"."DaysPastDue"' in finalized
    assert "STANDARD" in finalized and "SUBSTANDARD" in finalized
    assert validate_expression(finalized).valid


def test_sqlserver_catch_is_exception_handler_and_excluded_when_same_guard():
    sql = Path("samples/sql/01_NPA_Classification_Simple.sql").read_text()
    obj = split_objects(sql, "01.sql", detect_dialect(sql))[0]
    info = analyze_object(obj)
    sites = _assignment_sites(info, "COMPLETED")
    assert len(sites) == 2
    roles = [_infer_assignment_role(s.raw_sql) for s in sites]
    assert "EXCEPTION_HANDLER" in roles
    excluded = undeterminable_exception_sites(sites)
    assert len(excluded) == 1
    remaining = [s for s in sites if s not in excluded]
    composed = _compose_simple_assignment_expression(remaining, "RunStatus", "COMPLETED")
    assert composed
    assert 'THEN("Y")' in composed
    assert 'THEN("N")' not in composed


def test_decision_table_payload_keeps_display_expression():
    raw = """
    {
      "expression": "IF(\\"X\\"==1)THEN(\\"A\\")ELSE(\\"B\\")",
      "decision_table": {"decisionTableDetails": [{"derivedValue": "A", "sequenceNumber": 1}]}
    }
    """
    option, expression, dt_json, errors = _interpret_llm_output(raw)
    assert option == DerivationOption.DECISION_TABLE
    assert expression and "IF(" in expression
    assert dt_json and "decisionTableDetails" in dt_json
    assert not errors


def test_rewrite_keeps_entity_column_two_segment_paths():
    """`"FeeSchedule"."LateFee"` must not collapse to `"LateFee"`."""
    expr = '(COALESCE("LoanAccountCal"."LateFeeAmount", 0) + "FeeSchedule"."LateFee")'
    rewritten = rewrite_expression_to_platform_entities(
        expr,
        entity_name="FCT_LOAN_ACCOUNT",
        entity_name_map={
            "LoanAccountCal": "FCT_LOAN_ACCOUNT",
            "FeeSchedule": "FeeSchedule",
            "LateFee": "LateFee",  # polluted map entry must not win
        },
        alias_to_parts={"A": ("PRO", "LoanAccountCal"), "S": ("FeeSchedule",)},
    )
    assert '"FeeSchedule"."LateFee"' in rewritten
    assert rewritten.count('"LateFee"') == 1
    assert validate_expression(rewritten).valid


def test_datediff_and_dateadd_month_compose_from_samples():
    sql = Path("samples/sql/07_DPD_Bucket_Classification.sql").read_text()
    obj = split_objects(sql, "07.sql", detect_dialect(sql))[0]
    info = analyze_object(obj)
    sites = _assignment_sites(info, "DpdDays", target_table="LoanAccountCal")
    composed = _compose_simple_assignment_expression(sites, "FCT_LOAN_ACCOUNT", "DpdDays")
    assert composed and "DATEDIFF(" in composed
    assert validate_expression(composed).valid

    sql = Path("samples/sql/08_Loan_Restructuring_Eligibility.sql").read_text()
    obj = split_objects(sql, "08.sql", detect_dialect(sql))[0]
    info = analyze_object(obj)
    sites = _assignment_sites(info, "LastPaymentDueDate", target_table="LoanAccountCal")
    composed = _compose_simple_assignment_expression(sites, "FCT_LOAN_ACCOUNT", "LastPaymentDueDate")
    assert composed and "PERIOD(" in composed
    assert validate_expression(composed).valid


def test_numeric_column_addition_is_not_rejected_as_concat():
    result = validate_expression(
        '(COALESCE("FCT_LOAN_ACCOUNT"."LateFeeAmount", 0) + "FeeSchedule"."LateFee")'
    )
    assert result.valid


def test_in_subquery_and_sum_scalar_compose_from_samples():
    sql = Path("samples/sql/17_Dishonoured_Cheque_Penalty_Calc.sql").read_text()
    obj = split_objects(sql, "17.sql", detect_dialect(sql))[0]
    info = analyze_object(obj)
    sites = _assignment_sites(info, "ChequeBookSuspended", target_table="LoanAccountCal")
    composed = _compose_simple_assignment_expression(
        sites, "FCT_LOAN_ACCOUNT", "ChequeBookSuspended", procedure_sql=sql
    )
    assert composed and "DishonourCount" in composed and 'THEN("Y")' in composed
    assert validate_expression(composed).valid

    sql = Path("samples/sql/18_Batch_Job_Reconciliation.sql").read_text()
    obj = split_objects(sql, "18.sql", detect_dialect(sql))[0]
    info = analyze_object(obj)
    sites = _assignment_sites(info, "RetryCount", target_table="BatchFeedRegistry")
    composed = _compose_simple_assignment_expression(
        sites, "BatchFeedRegistry", "RetryCount", procedure_sql=sql
    )
    assert composed and "RETRY_QUEUED" in composed
    assert validate_expression(composed).valid

    sql = Path("samples/sql/16_Guarantee_Cover_Appropriation.sql").read_text()
    obj = split_objects(sql, "16.sql", detect_dialect(sql))[0]
    info = analyze_object(obj)
    sites = _assignment_sites(info, "AvailableBalance", target_table="GuaranteeFund")
    composed = _compose_simple_assignment_expression(
        sites, "GuaranteeFund", "AvailableBalance", procedure_sql=sql
    )
    assert composed and "SUM(" in composed
    assert validate_expression(composed).valid


def test_excel_export_matches_platform_columns(tmp_path):
    row = DDRow(
        entity_name="FCT_NPA_PRODUCT",
        column_name="AssetClass",
        column_type=ColumnType.PHYSICAL,
        derivation_option=DerivationOption.FORMULA_EXPRESSION,
        display_derivation_expression='IF("FCT_NPA_PRODUCT"."DaysPastDue"<=90)THEN("STANDARD")ELSE("LOSS")',
        effective_start_date=date(2026, 7, 20),
        status=DDStatus.ACTIVE,
        data_type="string",
        source_chain_id="c1",
    )
    out = export_dd_rows_excel([row], tmp_path / "dd_export.xlsx")
    assert out.exists()
    loaded = read_existing_dd_excel(out)
    assert len(loaded) == 1
    assert loaded[0]["entity_name"] == "FCT_NPA_PRODUCT"
    assert loaded[0]["column_name"] == "AssetClass"
    # Round-trip preserves the same schema headers as the sample Derivations export.
    from openpyxl import load_workbook

    wb = load_workbook(out)
    headers = [c.value for c in next(wb.active.iter_rows(min_row=1, max_row=1))]
    assert headers == COLUMNS


def test_sql_text_date_cast_map_to_platform_functions():
    import sqlglot
    from app.derivation.dd_generation_engine import _render_value_expression_to_4x

    cases = {
        "REPLACE(Name, 'a', 'b')": 'REPLACE(Name, "a", "b")',
        "CONVERT(VARCHAR(10), Amt)": 'CONVERT(Amt, "VARCHAR")',
        "DATEPART(year, BizDate)": 'DATEPART(BizDate, "year")',
        "LOWER(LTRIM(RTRIM(Code)))": 'LOWER(TRIM(TRIM(Code)))',
    }
    for sql, expected in cases.items():
        tree = sqlglot.parse_one(sql, read="tsql")
        rendered = _render_value_expression_to_4x(tree)
        assert rendered == expected, (sql, rendered)
        assert validate_expression(rendered).valid
